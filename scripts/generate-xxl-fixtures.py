#!/usr/bin/env python3
"""Convert the authored fixture matrix to genuine BIFF8 XLS and verify both files.

XLS authoring is the missing legacy-format capability; XLSX is authored by the
paired artifact-tool JavaScript builder. Requires xlwt 1.3.0, xlrd 2.0.1 and
openpyxl (read-only validation only). Use a temporary --target package directory
via PYTHONPATH rather than changing the bundled runtime installation.
"""
import hashlib
import json
from pathlib import Path
import sys

import openpyxl
import xlrd
import xlwt

manifest_path, output_dir = map(Path, sys.argv[1:3])
manifest = json.loads(manifest_path.read_text())
xls_path = output_dir / "模拟发布包" / "XXL-JOB配置.xls"
xlsx_path = output_dir / "模拟发布包" / "XXL-JOB配置.xlsx"
workbook = xlwt.Workbook(encoding="utf-8")
body_style = xlwt.easyxf("font: name Arial, height 220, colour black; align: wrap on, vert centre;")
title_style = xlwt.easyxf("font: name Arial, height 300, bold on, colour dark_blue;")
context_style = xlwt.easyxf("font: name Arial, height 220, colour gray50; align: vert centre;")
header_style = xlwt.easyxf("font: name Arial, height 220, bold on, colour white; pattern: pattern solid, fore_colour dark_blue; align: wrap on, vert centre, horiz centre;")
for spec in manifest["sheets"]:
    sheet = workbook.add_sheet(spec["name"])
    sheet.show_grid = False
    sheet.panes_frozen = True
    sheet.horz_split_pos = 5
    sheet.vert_split_pos = 1
    for column, width in enumerate(spec["widths"]):
        sheet.col(column).width = int(width / 7 * 256)
    for row, cells in enumerate(spec["matrix"]):
        sheet.row(row).height_mismatch = True
        sheet.row(row).height = int(spec["heights"][row] * 15)
        style = title_style if row == 1 else context_style if row == 2 else header_style if row == 4 else body_style
        for column, value in enumerate(cells):
            if value is not None:
                sheet.write(row, column, value, style)
workbook.save(str(xls_path))

assert xls_path.read_bytes()[:8] == bytes.fromhex("D0CF11E0A1B11AE1"), "XLS must be a real OLE/BIFF workbook"
assert xlsx_path.read_bytes()[:4] == b"PK\x03\x04", "XLSX must be an OOXML ZIP workbook"
legacy = xlrd.open_workbook(xls_path)
modern = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
checked_cells = 0
for spec in manifest["sheets"]:
    old_sheet = legacy.sheet_by_name(spec["name"])
    new_sheet = modern[spec["name"]]
    for row, cells in enumerate(spec["matrix"]):
        for column, expected in enumerate(cells):
            old_value = old_sheet.cell_value(row, column) if row < old_sheet.nrows and column < old_sheet.ncols else ""
            new_value = new_sheet.cell(row + 1, column + 1).value
            normalise = lambda value: "" if value is None else value
            assert normalise(old_value) == normalise(expected), (spec["name"], row, column, old_value, expected)
            assert normalise(new_value) == normalise(expected), (spec["name"], row, column, new_value, expected)
            checked_cells += 1
modern.close()
report = {
    "passed": True,
    "same_data_cells_checked": checked_cells,
    "sheets": [s["name"] for s in manifest["sheets"]],
    "signatures": {"xls": "D0CF11E0A1B11AE1 (OLE/BIFF8)", "xlsx": "504B0304 (OOXML ZIP)"},
    "files": [{"name": p.name, "bytes": p.stat().st_size, "sha256": hashlib.sha256(p.read_bytes()).hexdigest()} for p in [xlsx_path, xls_path]],
    "validation": "Full matrix equality via xlrd and openpyxl read-only, numeric values kept numeric; no model execution.",
}
(output_dir / "fixture-validation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
print(json.dumps(report, ensure_ascii=False))
